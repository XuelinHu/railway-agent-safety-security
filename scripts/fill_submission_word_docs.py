#!/usr/bin/env python3
"""Build the editable submission-side Word documents from verified metadata."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper" / "submission" / "submission-word"
TITLE = (
    "Provenance-Preserving Evidence-Gated Knowledge-Graph Augmentation "
    "for Auditable Entity-Relation Extraction"
)
JOURNAL = "Journal of Safety Science and Resilience"


def setup(document: Document, heading: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(heading)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)


def save(document: Document, name: str) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    document.save(OUT / name)


def highlights() -> None:
    document = Document()
    setup(document, "Highlights")
    for text in (
        "Separates graph-informed generation from deterministic evidence acceptance.",
        "Links accepted entities and relations to exact source spans and provenance.",
        "Improves relation F1 over matched source-only extraction on ADE and CoNLL04.",
        "Records auditable acceptance paths without asserting universal semantic truth.",
    ):
        document.add_paragraph(text, style="List Bullet")
    save(document, "Highlights.docx")


def title_page() -> None:
    document = Document()
    setup(document, "Title Page")
    document.add_heading("Article title", level=1)
    document.add_paragraph(TITLE)
    document.add_heading("Authors", level=1)
    document.add_paragraph(
        "Xuelin Hu¹; Xiaoqin Fu¹; Youjing Fu²; Jingchao Wang³,*; "
        "Ruishen Liu⁴; Charles Jumaa Katila⁵; Pengming Hu³"
    )
    document.add_heading("Affiliations", level=1)
    affiliations = (
        "¹ Liuzhou Railway Vocational Technical College, No. 2 Wenyuan Road, "
        "Yufeng District, Liuzhou 545000, Guangxi Zhuang Autonomous Region, China.",
        "² College of Earth and Environmental Sciences, Lanzhou University, "
        "Lanzhou 730000, China.",
        "³ School of Computer Science, Zhongyuan University of Technology, "
        "Zhengzhou 450007, China.",
        "⁴ School of Artificial Intelligence, Luoyang Normal University, "
        "Luoyang 471000, China.",
        "⁵ School of Computing and Mathematics, The Co-operative University of Kenya, Kenya.",
    )
    for item in affiliations:
        document.add_paragraph(item)
    document.add_heading("Corresponding author", level=1)
    document.add_paragraph("Jingchao Wang; email: static@zut.edu.cn")
    document.add_heading("Author email addresses", level=1)
    for item in (
        "Xuelin Hu: huxuelinai@gmail.com; huxl@ltzy.edu.cn",
        "Xiaoqin Fu: xiaoqin.fu@qq.com; fuxiaoqin@ltzy.edu.cn",
        "Youjing Fu: fuyj2025@lzu.edu.cn",
        "Jingchao Wang: static@zut.edu.cn",
        "Ruishen Liu: liuruishen@lynu.edu.cn",
        "Charles Jumaa Katila: ckatila@cuk.ac.ke",
        "Pengming Hu: hupengming@zut.edu.cn",
    ):
        document.add_paragraph(item)
    document.add_heading("ORCID identifiers", level=1)
    document.add_paragraph(
        "Xuelin Hu: 0000-0002-4475-3034; Xiaoqin Fu: 0009-0003-5123-8393; "
        "Youjing Fu: 0009-0004-1369-7879."
    )
    document.add_heading("Funding", level=1)
    document.add_paragraph(
        "This work was supported by the Science and Technology Research Project of "
        "Henan Province (grant no. 262102211055) and the Key Scientific Research "
        "Project of Colleges and Universities in Henan Province (grant no. 27AQ520018)."
    )
    save(document, "Title Page.docx")


def author_statement() -> None:
    document = Document()
    setup(document, "CRediT Author Statement")
    document.add_paragraph(TITLE)
    document.add_paragraph(
        "The author names below have been synchronized with the title page. Individual "
        "CRediT roles require confirmation by all authors before submission; no roles "
        "have been inferred from author order."
    )
    for name in (
        "Xuelin Hu", "Xiaoqin Fu", "Youjing Fu", "Jingchao Wang",
        "Ruishen Liu", "Charles Jumaa Katila", "Pengming Hu",
    ):
        document.add_paragraph(f"{name}: CRediT role(s) to be confirmed by the authors.")
    save(document, "author statement.docx")


def cover_letter() -> None:
    document = Document()
    setup(document, "Cover Letter")
    document.add_paragraph(date.today().strftime("%B %d, %Y"))
    document.add_paragraph("Dear Editors,")
    document.add_paragraph(
        f'We submit the manuscript entitled "{TITLE}" for consideration in {JOURNAL}.'
    )
    document.add_paragraph(
        "The manuscript addresses a central limitation of graph-augmented entity-relation "
        "extraction: retrieved associations and fluent generated triples are not necessarily "
        "facts stated in the current source sentence. We separate graph-informed candidate "
        "generation from deterministic acceptance, reconstruct exact source spans, enforce "
        "type-compatible relations, and retain an auditable provenance record. Evaluation "
        "on ADE and CoNLL04 shows higher relation F1 than the matched source-only generator "
        "on both benchmarks, while the manuscript explicitly limits the gates to structural "
        "admissibility and source linkage rather than semantic or clinical truth."
    )
    document.add_paragraph(
        "This work is original, has not been published previously, and is not under "
        "consideration by another journal. All authors must approve the final submitted "
        "version and the accompanying declarations."
    )
    document.add_paragraph("Thank you for considering our manuscript.")
    document.add_paragraph("Sincerely,")
    document.add_paragraph(
        "Jingchao Wang\nCorresponding author\nSchool of Computer Science, "
        "Zhongyuan University of Technology\nZhengzhou 450007, China\n"
        "Email: static@zut.edu.cn"
    )
    save(document, "cover letter.docx")


def competing_interests() -> None:
    document = Document()
    setup(document, "Declaration of Competing Interests")
    document.add_paragraph(TITLE)
    document.add_paragraph(
        "The authors declare that they have no known competing financial interests or "
        "personal relationships that could have appeared to influence the work reported "
        "in this paper."
    )
    document.add_paragraph(
        "This declaration must be confirmed by every author before submission."
    )
    save(document, "declaration-of-competing-interests.docx")


def main() -> None:
    highlights()
    title_page()
    author_statement()
    cover_letter()
    competing_interests()


if __name__ == "__main__":
    main()
